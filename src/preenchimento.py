from docx import Document


def _substituir_em_runs(paragrafo, antigo, novo):
    inicio_busca = 0
    while True:
        runs = paragrafo.runs
        textos = [r.text for r in runs]
        completo = "".join(textos)
        idx = completo.find(antigo, inicio_busca)
        if idx == -1:
            return
        fim = idx + len(antigo)

        pos = 0
        run_ini = run_fim = None
        for i, t in enumerate(textos):
            if run_ini is None and idx < pos + len(t):
                run_ini, off_ini = i, idx - pos
            if run_ini is not None and fim <= pos + len(t):
                run_fim, off_fim = i, fim - pos
                break
            pos += len(t)

        if run_ini == run_fim:
            t = textos[run_ini]
            runs[run_ini].text = t[:off_ini] + novo + t[off_fim:]
        else:
            runs[run_ini].text = textos[run_ini][:off_ini] + novo
            for i in range(run_ini + 1, run_fim):
                runs[i].text = ""
            runs[run_fim].text = textos[run_fim][off_fim:]

        inicio_busca = idx + len(novo)


def _todos_paragrafos(documento):
    def de_tabelas(tabelas):
        for tabela in tabelas:
            for linha in tabela.rows:
                for celula in linha.cells:
                    yield from celula.paragraphs
                    yield from de_tabelas(celula.tables)

    yield from documento.paragraphs
    yield from de_tabelas(documento.tables)
    for secao in documento.sections:
        for parte in (secao.header, secao.footer):
            yield from parte.paragraphs
            yield from de_tabelas(parte.tables)


LIMPEZAS = [(", , ", ", "), (",  ,", ","), (" ,", ",")]


def preencher_modelo(caminho_modelo, caminho_saida, substituicoes, remover_destaque=True):
    documento = Document(caminho_modelo)

    for paragrafo in _todos_paragrafos(documento):
        if "{{" in paragrafo.text:
            for marcador, valor in substituicoes.items():
                if marcador in paragrafo.text:
                    _substituir_em_runs(paragrafo, marcador, str(valor))
            for antigo, novo in LIMPEZAS:
                if antigo in paragrafo.text:
                    _substituir_em_runs(paragrafo, antigo, novo)

        if remover_destaque:
            for run in paragrafo.runs:
                run.font.highlight_color = None

    documento.save(caminho_saida)


def marcadores_restantes(caminho_docx):
    import re
    documento = Document(caminho_docx)
    achados = set()
    for p in _todos_paragrafos(documento):
        achados.update(re.findall(r"\{\{[A-Z_]+\}\}", p.text))
    return sorted(achados)
