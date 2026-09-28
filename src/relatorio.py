from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AZUL = PatternFill("solid", fgColor="1F4E79")
VERDE = PatternFill("solid", fgColor="E2EFDA")
AMARELO = PatternFill("solid", fgColor="FFF2CC")
BRANCO_NEGRITO = Font(bold=True, color="FFFFFF")


def _tabela(planilha, cabecalho, linhas, larguras):
    planilha.append(cabecalho)
    for celula in planilha[1]:
        celula.fill = AZUL
        celula.font = BRANCO_NEGRITO
        celula.alignment = Alignment(vertical="center")
    for linha in linhas:
        planilha.append(linha)
    for i, largura in enumerate(larguras, 1):
        planilha.column_dimensions[get_column_letter(i)].width = largura
    planilha.freeze_panes = "A2"
    if linhas:
        planilha.auto_filter.ref = planilha.dimensions


def salvar_relatorio(caminho, resumo, registros, avisos_gerais):
    livro = Workbook()

    aba = livro.active
    aba.title = "Resumo"
    aba.append(["Relatório de geração de declarações"])
    aba["A1"].font = Font(bold=True, size=14)
    aba.append([])
    for chave, valor in resumo.items():
        aba.append([chave, valor])
        aba.cell(row=aba.max_row, column=1).font = Font(bold=True)
    if avisos_gerais:
        aba.append([])
        aba.append(["Avisos gerais"])
        aba.cell(row=aba.max_row, column=1).font = Font(bold=True)
        for aviso in avisos_gerais:
            aba.append(["", aviso])
    aba.column_dimensions["A"].width = 28
    aba.column_dimensions["B"].width = 70

    cabecalho = ["Linha na planilha", "Empreendimento", "CNPJ", "Município", "Arquivo gerado",
                 "Situação", "Observações"]
    larguras = [16, 42, 22, 20, 46, 22, 60]

    todas = livro.create_sheet("Declarações")
    _tabela(todas, cabecalho, [list(r.values()) for r in registros], larguras)
    for linha in todas.iter_rows(min_row=2, min_col=6, max_col=6):
        for celula in linha:
            celula.fill = VERDE if celula.value == "OK" else AMARELO

    pendentes = [r for r in registros if r["Situação"] != "OK"]
    aba_pend = livro.create_sheet("Pendências")
    if pendentes:
        _tabela(aba_pend, cabecalho, [list(r.values()) for r in pendentes], larguras)
    else:
        aba_pend.append(["Nenhuma pendência. Todas as declarações foram geradas com os dados completos."])
        aba_pend.column_dimensions["A"].width = 90

    livro.save(caminho)